from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("交付文件/生活服务AI小程序_PRD_v0.4_内容修订红字.docx")
OUTPUT = Path("交付文件/生活服务AI小程序_PRD_v0.4_最终版.docx")

FONT = "Hiragino Sans GB"
INK = RGBColor(0x24, 0x2B, 0x2E)
MUTED = RGBColor(0x67, 0x73, 0x78)
TEAL = RGBColor(0x24, 0x4E, 0x5F)
TEAL_LIGHT = RGBColor(0x37, 0x69, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "315D70"
ROW_FILL = "F3F7F8"


def set_east_asian_font(font, name):
    font.name = name
    rpr = font._element
    rfonts = getattr(rpr, "rFonts", None)
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def set_run(run, size, color, bold=None, italic=None):
    set_east_asian_font(run.font, FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def para_style(paragraph):
    name = paragraph.style.name if paragraph.style else "Normal"
    if name == "Title":
        return 25, TEAL, 0, 9, 1.0, True
    if name == "Heading 1":
        return 16, TEAL, 17, 8, 1.1, True
    if name == "Heading 2":
        return 13, TEAL_LIGHT, 12, 5, 1.1, True
    if name == "Heading 3":
        return 11.5, TEAL_LIGHT, 9, 4, 1.1, True
    if name.startswith("List"):
        return 10.5, INK, 0, 3, 1.18, None
    return 10.5, INK, 0, 5, 1.2, None


def apply_paragraph(paragraph, in_table=False):
    size, color, before, after, spacing, bold = para_style(paragraph)
    if in_table:
        size, color, before, after, spacing = 9.2, INK, 0, 2, 1.12
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    if paragraph.style and paragraph.style.name.startswith("Heading"):
        pf.keep_with_next = True
        pf.keep_together = True
    for run in paragraph.runs:
        set_run(run, size, color, bold=bold)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


doc = Document(SOURCE)

# Normalize document-wide style definitions first.
for style_name, size, color, before, after, line, bold in [
    ("Normal", 10.5, INK, 0, 5, 1.2, False),
    ("Title", 25, TEAL, 0, 9, 1.0, True),
    ("Heading 1", 16, TEAL, 17, 8, 1.1, True),
    ("Heading 2", 13, TEAL_LIGHT, 12, 5, 1.1, True),
    ("Heading 3", 11.5, TEAL_LIGHT, 9, 4, 1.1, True),
    ("List Paragraph", 10.5, INK, 0, 3, 1.18, False),
]:
    style = doc.styles[style_name]
    set_east_asian_font(style.font, FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    if style_name != "List Paragraph":
        style.font.bold = bold
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line

# Normalize all document paragraphs while retaining the existing hierarchy, alignment and numbering.
for paragraph in doc.paragraphs:
    apply_paragraph(paragraph)
    if paragraph.style and paragraph.style.name == "Title":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Normalize every table and remove revision red while keeping readable header contrast.
for table in doc.tables:
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            header = row_index == 0
            set_cell_shading(cell, HEADER_FILL if header else (ROW_FILL if row_index % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                apply_paragraph(paragraph, in_table=True)
                if header:
                    for run in paragraph.runs:
                        run.font.color.rgb = WHITE
                        run.bold = True

# Reapply consistent fonts/colors in section headers and any text inside tables that bypassed paragraph styles.
for paragraph in doc.paragraphs:
    apply_paragraph(paragraph)
for table in doc.tables:
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                apply_paragraph(paragraph, in_table=True)
                if ri == 0:
                    for run in paragraph.runs:
                        run.font.color.rgb = WHITE
                        run.bold = True

doc.save(OUTPUT)
print(OUTPUT)
