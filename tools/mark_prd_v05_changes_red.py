from pathlib import Path

from docx import Document
from docx.shared import RGBColor


SOURCE = Path("交付文件/生活服务AI小程序_PRD_v0.5.docx")
OUTPUT = Path("交付文件/生活服务AI小程序_PRD_v0.5_红色修订.docx")
RED = RGBColor(0xC0, 0x00, 0x00)


def color_paragraph(paragraph):
    for run in paragraph.runs:
        if run.text:
            run.font.color.rgb = RED


def has_any(text, phrases):
    return any(phrase in text for phrase in phrases)


doc = Document(SOURCE)

# Paragraphs added or rewritten in this revision.
paragraph_phrases = [
    "版本：V0.5",
    "城市边界：",
    "首期仅提供提醒偏好设置与计划到期状态展示",
    "用户侧轻量反馈作为 P1 能力",
    "如用户未说明目的地，系统先确认目的地是否为上海",
    "用户可在回答中查看推荐区域及参考依据",
    "更新计划页中相关事项的到期状态展示",
    "配置提醒偏好。提醒偏好仅影响本地展示",
    "首次对话优先收集会直接影响计划的入职时间",
    "硬截止仅能来自用户明确给出的截止日期",
    "仅删除本地资料、聊天记录和提醒偏好，不删除已确认日程",
    "5.7 F06 回答反馈（P1）",
    "在每次助手回答下提供“有帮助 / 不准确 / 不相关”",
    "6.5 日期依赖与硬截止规则",
    "搬家/入住日期变更时，系统检查未完成的交接",
    "8.3.1 知识维护与失效处理",
    "每条可引用的政策、入口和区域数据应记录维护责任",
    "若部署版本接入外部大模型服务",
    "删除浏览器内保存的资料、聊天记录和提醒偏好，已确认日程保留",
]

for paragraph in doc.paragraphs:
    if has_any(paragraph.text, paragraph_phrases):
        color_paragraph(paragraph)

# Only the table cells altered in this revision are colored.
changed_cells = {
    (0, 6, 0),
    (3, 2, 2),
    (4, 5, 1), (4, 5, 2),
    (5, 3, 2),
    (8, 3, 1), (8, 3, 2), (8, 4, 1), (8, 4, 2),
    (15, 6, 1),
    (16, 1, 1), (16, 2, 1), (16, 4, 1), (16, 6, 1),
    (16, 7, 0), (16, 7, 1), (16, 7, 2),
    (16, 8, 0), (16, 8, 1), (16, 8, 2),
}

for table_index, row_index, col_index in changed_cells:
    cell = doc.tables[table_index].cell(row_index, col_index)
    for paragraph in cell.paragraphs:
        color_paragraph(paragraph)

doc.save(OUTPUT)
print(OUTPUT)
