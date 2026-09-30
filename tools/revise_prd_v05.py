from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE = Path("/Users/sususu/Desktop/桌面 - Sususu's MacBook/国内简历/AI产品经理/课/生活助手小程序/生活服务AI小程序_PRD_v0.4.docx")
OUTPUT = Path("交付文件/生活服务AI小程序_PRD_v0.5.docx")


def set_paragraph(paragraph, text):
    """Replace text while retaining the paragraph's existing character style."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_text(paragraphs, old, new, required=True):
    found = False
    for paragraph in paragraphs:
        if old in paragraph.text:
            set_paragraph(paragraph, paragraph.text.replace(old, new))
            found = True
    if required and not found:
        raise RuntimeError(f"Missing paragraph text: {old}")


def cell_set(cell, text):
    paragraph = cell.paragraphs[0]
    set_paragraph(paragraph, text)
    for extra in cell.paragraphs[1:]:
        set_paragraph(extra, "")


def insert_after(paragraph, text, style_source):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style_source.style
    if style_source.runs:
        run = new_para.add_run(text)
        source_rpr = style_source.runs[0]._element.rPr
        if source_rpr is not None:
            run._element.insert(0, deepcopy(source_rpr))
    else:
        new_para.add_run(text)
    return new_para


doc = Document(SOURCE)
p = doc.paragraphs

# Version and scope.
set_paragraph(p[1], "版本：V0.5")
set_paragraph(
    p[25],
    "城市边界：用户可来自任意城市，但首期仅对“目的地为上海”的就业搬迁提供结构化计划、上海知识与入口支持。目的地非上海时，仅提供通用事项清单和官方核验路径，不输出当地政策、区域或费用结论。",
)
replace_text(
    p,
    "5. 提醒的设置、授权、发送结果可解释，不让用户误以为保存计划就一定会收到微信通知。",
    "5. 首期仅提供提醒偏好设置与计划到期状态展示；不让用户误以为保存计划就一定会收到站内或微信通知。",
)
replace_text(
    p,
    "6. 建立 AI 回答质量评测和用户反馈机制，为后续迭代提供依据。",
    "6. 建立 AI 回答质量评测机制；用户侧轻量反馈作为 P1 能力，为后续迭代提供依据。",
)

# Scenarios and explicit product boundaries.
replace_text(p, "系统补问目标城市、办公位置、入职日期、计划入住日期和首次安家可用资金。", "如用户未说明目的地，系统先确认目的地是否为上海；在上海范围内，再补问办公位置、入职日期、计划入住日期和首次安家可用资金。")
replace_text(p, "用户可以收藏感兴趣的区域，再自行通过外部渠道找房。", "用户可在回答中查看推荐区域及参考依据，再自行通过外部渠道找房。")
replace_text(p, "用户确认后保存新计划并更新内部提醒。", "用户确认后保存新计划，并更新计划页中相关事项的到期状态展示。")

# Architecture and state wording.
replace_text(p, "用户可在日历标记完成、删除自定义日程，或在“我的”页面更正资料、查看历史输入、配置站内提醒。提醒设置不等同于微信通知已发送。", "用户可在日历标记完成、删除自定义日程，或在“我的”页面更正资料、查看历史输入、配置提醒偏好。提醒偏好仅影响本地展示，不等同于站内或微信通知已发送。")

# F01/F02/F04.
replace_text(
    p,
    "首页提示用户优先提供入职时间、公司位置、通勤时长、搬家完成时间、一次性搬家预算和是否接受合租；需要区域建议时再补充月租预算。",
    "首次对话优先收集会直接影响计划的入职时间、公司位置和搬家/入住节点，最多追问 3 项；通勤时长、一次性搬家预算、月租预算和是否接受合租仅在区域推荐、预算判断或相关计划需要时补充。",
)
replace_text(
    p,
    "新增：生成标题、开始/结束日期、是否硬截止及来源标识；日期未知时进入“日期待确认”，不伪造日期。",
    "新增：生成标题、开始/结束日期、是否硬截止及来源标识。硬截止仅能来自用户明确给出的截止日期，或已核验的官方/单位规则；日期未知时进入“日期待确认”，不伪造日期。",
)
replace_text(
    p,
    "页面提供“清空我的资料和聊天记录”，用于删除本地用户状态。",
    "页面提供“清空我的资料和聊天记录”，仅删除本地资料、聊天记录和提醒偏好，不删除已确认日程；若后续提供“清空全部本地数据”，必须单独列明会删除日程并二次确认。",
)

# Knowledge, privacy, dependency rules, and feedback requirement.
anchor = p[117]
heading = insert_after(anchor, "8.3.1 知识维护与失效处理", p[113])
insert_after(
    heading,
    "每条可引用的政策、入口和区域数据应记录维护责任、核验日期、适用城市、来源 URL 与有效状态。入口失效、内容冲突或超过复核周期时，应降级为“待核验”，不得继续作为确定性结论或正式计划依据；产品向用户展示官方查找方向与最近核验时间。",
    p[114],
)
replace_text(
    p,
    "资料、聊天和日程默认保存于浏览器本地；发送问答时只携带与当前问题相关的资料与近期对话。",
    "资料、聊天和日程默认保存于浏览器本地；发送问答时，前端仅向本地规划服务携带与当前问题相关的资料与近期对话。若部署版本接入外部大模型服务，应在用户可见的隐私说明中明确数据会发送给该服务、发送字段范围及其适用的数据处理规则。",
)
replace_text(
    p,
    "用户可通过“清空我的资料和聊天记录”删除浏览器内保存的状态；跨设备、服务端保留和备份策略需在后续账号化前另行定义。",
    "用户可通过“清空我的资料和聊天记录”删除浏览器内保存的资料、聊天记录和提醒偏好，已确认日程保留；跨设备、服务端保留和备份策略需在后续账号化前另行定义。",
)

anchor = p[95]
heading = insert_after(anchor, "6.5 日期依赖与硬截止规则", p[90])
insert_after(
    heading,
    "搬家/入住日期变更时，系统检查未完成的交接、通勤试走、宽带核验等依赖事项，并仅在 pending 中展示受影响变化；已完成事项保留记录。用户明确的入职、预约或材料截止日期，以及经核验的官方/单位截止规则，标记为硬截止。硬截止不得自动顺延、删除或降级；发生冲突时，系统只提示风险、要求再次核验，并由用户确认后保存。每项依赖关系应注明触发事项、受影响事项、处理方式与依据。",
    p[91],
)

anchor = p[80]
heading = insert_after(anchor, "5.7 F06 回答反馈（P1）", p[77])
insert_after(
    heading,
    "在每次助手回答下提供“有帮助 / 不准确 / 不相关”三个轻量反馈入口，并允许用户选填原因。反馈不改变日程或资料，不收集证件、账号等敏感信息；产品按场景、回答模式和来源状态汇总，用于人工复核和提示词、知识维护迭代。首期原型未交付该入口，不以负反馈率作为已实测指标。",
    p[78],
)

# Tables: retain all non-affected content and change only impacted cells.
t = doc.tables
cell_set(t[0].cell(6, 0), "提醒偏好设置、计划到期状态展示")
cell_set(t[3].cell(2, 2), "默认仅保存在当前浏览器；清空资料与聊天记录仅移除资料、聊天记录和提醒偏好，不删除已确认日程。")
cell_set(t[4].cell(5, 1), "未设置 / 已保存 / 待接入")
cell_set(t[4].cell(5, 2), "当前前端可保存提醒偏好与展示到期状态；未接入站内推送、微信授权或真实发送能力。")
cell_set(t[5].cell(3, 2), "资料维护、历史输入、本地清除、提醒偏好与到期状态展示")
cell_set(t[8].cell(3, 1), "可保存开关、提前 3 天和到期当天的显示时间。")
cell_set(t[8].cell(3, 2), "仅为本地偏好与到期状态展示，不代表后台推送已实现。")
cell_set(t[8].cell(4, 1), "后续接入能力：微信订阅提醒")
cell_set(t[8].cell(4, 2), "需另行定义授权、模板、发送、失败、重试与取消订阅流程；首期不展示为已交付能力。")
cell_set(t[15].cell(6, 1), "清空资料与聊天记录仅删除本地资料、聊天记录和提醒偏好，已确认日程保留；另设“清空全部本地数据”时必须单独确认并明确包含日程。页面提醒用户不要输入敏感信息。")
cell_set(t[16].cell(1, 1), "任务是否含明确动作、时间依据、前置与完成标准；每次版本验收随机抽取不少于 30 条核心场景回答，由产品与知识维护人员按统一标注表复核。")
cell_set(t[16].cell(2, 1), "确定性政策/入口是否与适用证据匹配。严重错误包括虚构政策、网址、截止日期、预约状态，或将待核验/样本信息表述为确定事实。")
cell_set(t[16].cell(4, 1), "改期后受影响的未完成事项是否正确呈现；核心用例须明确覆盖搬家、交接、通勤试走、宽带核验和硬截止冲突。")
cell_set(t[16].cell(6, 1), "“没帮助/不准确/不相关”等用户反馈占回答反馈总数的比例；P1 反馈入口上线前，仅记录人工测试反馈，不作为线上指标。")

# Add product funnel metrics by expanding the existing metrics table with two rows.
for values in [
    ("计划确认率", "产生 pending 提议后，用户确认写入 confirmed 的会话占比。", "灰度阶段按场景观察，不设置未经验证的绝对目标。"),
    ("关键事项推进率", "已确认计划中，用户主动标记完成至少一项关键任务的用户占比。", "作为产品内推进指标；不以用户是否完成外部办理或搬迁作为首期硬性门槛。"),
]:
    cells = t[16].add_row().cells
    for cell, value in zip(cells, values):
        cell_set(cell, value)

doc.save(OUTPUT)
print(OUTPUT)
