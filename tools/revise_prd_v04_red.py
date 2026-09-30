from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph


SOURCE = Path("/Users/sususu/Desktop/桌面 - Sususu's MacBook/国内简历/AI产品经理/课/生活助手小程序/生活服务AI小程序_PRD_v0.4.docx")
OUTPUT = Path("交付文件/生活服务AI小程序_PRD_v0.4_内容修订红字.docx")
RED = RGBColor(0xC0, 0x00, 0x00)


def locate(text):
    matches = [p for p in doc.paragraphs if p.text.strip() == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph for {text!r}, found {len(matches)}")
    return matches[0]


def locate_contains(fragment):
    matches = [p for p in doc.paragraphs if fragment in p.text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph containing {fragment!r}, found {len(matches)}")
    return matches[0]


def set_paragraph(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        paragraph.runs[0].font.color.rgb = RED
    else:
        run = paragraph.add_run(text)
        run.font.color.rgb = RED


def insert_after(anchor, text, template, color=True):
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = template.style
    if template._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(template._p.pPr))
    run = paragraph.add_run(text)
    if template.runs and template.runs[0]._element.rPr is not None:
        run._element.insert(0, deepcopy(template.runs[0]._element.rPr))
    if color:
        run.font.color.rgb = RED
    return paragraph


def set_cell(cell, text, color=True):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
        if color:
            p.runs[0].font.color.rgb = RED
    else:
        r = p.add_run(text)
        if color:
            r.font.color.rgb = RED
    for extra in cell.paragraphs[1:]:
        for r in extra.runs:
            r.text = ""


doc = Document(SOURCE)

# Clarify the first-release city boundary in the core scenario.
old = locate("系统补问目标城市、办公位置、入职日期、计划入住日期和首次安家可用资金。先输出建议完成租房的时间节点、租房注意事项、符合条件的区域参考、生活费用预估，以及可并行的入职要求确认任务；没有正式住址时，不直接生成确定日期的宽带安装安排。")
set_paragraph(old, "系统先确认目的城市是否为上海；首期仅对目的地为上海的就业搬迁提供结构化计划、上海事项与区域参考。目的地不是上海时，说明当前支持范围，仅提供通用准备清单和官方核验方向，不输出当地政策、区域租金、通勤或办理入口结论。目的地为上海时，再补问办公位置、入职日期、计划入住日期和首次安家可用资金，并输出租房节点、注意事项、区域参考、生活费用估算及可并行的入职要求核验任务；没有正式住址时，不生成确定日期的宽带安装安排。")

# Replace the unsupported area-favorites promise while preserving the housing boundary.
old = locate("用户可以收藏感兴趣的区域，再自行通过外部渠道找房。小程序不展示、导入、推荐或比较具体房源，也不根据用户发来的房源链接评估哪一套更好。")
set_paragraph(old, "用户可在回答中查看推荐区域及参考依据，再自行通过外部渠道找房。小程序不展示、导入、推荐或比较具体房源，也不根据用户发来的房源链接评估哪一套更好。")

# Align all statements about clearing local data: this action keeps confirmed calendar items.
old = locate("历史输入仅展示在当前浏览器中保存的用户输入。页面提供“清空我的资料和聊天记录”，用于删除本地用户状态。")
set_paragraph(old, "历史输入仅展示在当前浏览器中保存的用户输入。页面提供“清空我的资料和聊天记录”，仅删除本地资料、聊天记录和提醒偏好，保留已确认计划日历。")
old = locate("4. 用户可通过“清空我的资料和聊天记录”一次性删除本地资料和历史输入；该操作不删除计划日历。")
set_paragraph(old, "4. 用户可通过“清空我的资料和聊天记录”一次性删除本地资料、历史输入和提醒偏好；该操作保留已确认计划日历。")
set_cell(doc.tables[8].cell(6, 1), "“清空我的资料和聊天记录”删除本地资料、聊天记录和提醒偏好，保留已确认日程；页面须明确提示清除范围，并提醒用户不要输入敏感信息。")
old = locate("用户可通过“清空我的资料和聊天记录”删除浏览器内保存的状态；跨设备、服务端保留和备份策略需在后续账号化前另行定义。")
set_paragraph(old, "用户可通过“清空我的资料和聊天记录”删除浏览器内保存的资料、聊天记录和提醒偏好；已确认日程保留。跨设备、服务端保留和备份策略需在后续账号化前另行定义。")

# Repair incomplete and corrupted rules in section 6.4.
old = locate("普通建议不得产生；仅计划生成或明确日程变更可产生日程提议。")
set_paragraph(old, "普通建议不得产生 operations；仅计划生成或明确日程变更可产生日程提议。")
old = locate("系统以计划清单为下一轮提议基准；如果用户撤回未确认提议，前端清除 计划清单而不影响 计划清单确认计划日历。")
set_paragraph(old, "系统以 confirmed 已确认日程作为当前有效计划，以 pending 待确认提议作为下一轮变更草稿；用户撤回待确认提议时，前端清除 pending，不影响 confirmed 日程。")

# Add a concrete knowledge maintenance and expiry rule after the existing source requirements.
anchor = locate("区域、路线、费用和配套数据应展示口径与查询/核验时间；未知费用不能按零处理。")
knowledge_heading = insert_after(anchor, "8.3.1 知识维护与失效处理", locate("8.3 检索与展示要求"))
insert_after(knowledge_heading, "每类政策、办事入口和区域数据均须指定维护角色；知识记录保留来源 URL、适用地区、适用对象、最近核验日期和下次复核日期。政策与官方入口至少每 30 天复核一次；区域租金、路线和配套数据按数据提供方更新周期复核，并在展示时标明样本/估算口径及更新时间。链接失效、内容冲突或超过复核日期时，立即降级为“待核验”，暂停作为确定性结论或明确匹配依据；完成复核后方可恢复。", locate("每项政策性结论应能对应具体证据；来源不足时降低确定性，不用无关链接填充。"))

# Add the promised P1 feedback feature before section 6, and list it in the feature overview.
anchor = locate("清除浏览器站点数据也会删除这些记录。")
feedback_heading = insert_after(anchor, "5.8 F06 回答反馈（P1）", locate("5.7 历史记录模块"))
feedback_intro = insert_after(feedback_heading, "功能说明：", locate("功能说明："))
feedback_body = insert_after(feedback_intro, "每条助手回答下提供“有帮助 / 不准确 / 不相关”三个反馈选项，用户可选填原因。反馈仅用于回答质量复核，不直接修改资料、日程或提醒设置；汇总时按场景、回答模式和来源状态分类，供产品与知识维护迭代。该入口属于 P1，未上线前不统计线上负反馈率。", locate_contains("历史输入仅用于回顾，不支持恢复整段历史会话，也不支持逐条删除。"))

feature_table = doc.tables[3]
row = feature_table.add_row()
set_cell(row.cells[0], "P1")
set_cell(row.cells[1], "回答反馈")
set_cell(row.cells[2], "回答有帮助/不准确/不相关反馈及选填原因；用于质量复核，不改变资料或日程。")

# Make the acceptance and AI quality metrics reproducible.
metrics = doc.tables[9]
set_cell(metrics.cell(1, 1), "每次版本验收抽取不少于 30 条覆盖核心场景的回答，由 2 名评审按同一评分表独立判断任务是否含明确动作、时间依据、前置条件和完成标准；分歧复核后计算通过率。")
set_cell(metrics.cell(2, 1), "确定性政策、资格条件、截止日期和办理入口是否与适用证据匹配。严重错误包括虚构政策/网址/截止日期、把待核验或样本信息说成确定事实、误导用户错过硬截止，或声称未确认的日程变更已生效。")
set_cell(metrics.cell(5, 1), "编造规则、价格、通勤、网址、预约状态或保证结果的比例；严重错误按“来源准确性”指标中的定义判定。")
set_cell(metrics.cell(6, 1), "反馈率 = 收到反馈的回答数 ÷ 已展示回答数；同时按有帮助、不准确、不相关和用户选填原因分组。")
set_cell(metrics.cell(6, 2), "反馈入口上线后按场景跟踪；入口上线前只记录内部测试反馈，不作为线上负反馈指标。")

doc.save(OUTPUT)
print(OUTPUT)
